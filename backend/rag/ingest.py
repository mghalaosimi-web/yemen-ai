from pathlib import Path
import json, csv, re
from .chunker import chunk_text

TEXT_EXTENSIONS = {".txt", ".md", ".csv", ".json", ".py", ".html", ".htm", ".pdf", ".docx"}


def _clean(text: str) -> str:
    text = str(text or "").replace("\x00", " ")
    text = re.sub(r"[ \t]+", " ", text)
    text = re.sub(r"\n{3,}", "\n\n", text)
    return text.strip()


def _ocr_pdf_page(page, page_number: int):
    """Best-effort OCR for scanned PDFs. Optional runtime dependencies fail safely."""
    try:
        import pytesseract
        from PIL import Image
        pix = page.get_pixmap(matrix=__import__('fitz').Matrix(2, 2), alpha=False)
        image = Image.frombytes("RGB", [pix.width, pix.height], pix.samples)
        # Arabic + English supports mixed university and technical documents.
        text = pytesseract.image_to_string(image, lang='Arabic+eng')
        return _clean(text), True
    except Exception:
        return "", False


def read_document_pages(path: str):
    p = Path(path); suffix = p.suffix.lower()
    if suffix not in TEXT_EXTENSIONS:
        raise ValueError(f"Unsupported ingestion type: {suffix}")

    if suffix == '.pdf':
        # PyMuPDF is preferred because it gives stable page-level extraction and rendering.
        try:
            import fitz
            doc = fitz.open(str(p)); pages=[]; scanned=[]
            for i, page in enumerate(doc, 1):
                text = _clean(page.get_text('text'))
                ocr_used = False
                if len(re.sub(r'\s+', '', text)) < 20:
                    text, ocr_used = _ocr_pdf_page(page, i)
                if text:
                    pages.append((i, text, {'ocr': ocr_used}))
                elif ocr_used:
                    scanned.append(i)
            if not pages:
                raise ValueError('PDF contains no readable text. OCR could not extract usable content from the scanned pages.')
            return pages
        except ValueError:
            raise
        except Exception:
            # Conservative fallback when PyMuPDF is unavailable.
            try:
                from pypdf import PdfReader
                reader=PdfReader(str(p)); pages=[]
                for i,page in enumerate(reader.pages,1):
                    text=_clean(page.extract_text() or '')
                    if text: pages.append((i,text,{'ocr':False}))
                if not pages: raise ValueError('PDF contains no extractable text. Install OCR support for scanned PDFs.')
                return pages
            except ValueError: raise
            except Exception as e: raise ValueError(f'Could not read PDF: {e}')

    if suffix == '.docx':
        try:
            from docx import Document
            doc=Document(str(p)); parts=[]
            for para in doc.paragraphs:
                if para.text.strip(): parts.append(para.text)
            for table in doc.tables:
                for row in table.rows:
                    cells=[c.text.strip() for c in row.cells]
                    if any(cells): parts.append(' | '.join(cells))
            text=_clean('\n'.join(parts))
            if not text: raise ValueError('DOCX contains no readable text')
            return [(1,text,{'ocr':False})]
        except ValueError: raise
        except Exception as e: raise ValueError(f'Could not read DOCX: {e}')

    raw=p.read_text(encoding='utf-8',errors='ignore')
    if suffix=='.json':
        try: raw=json.dumps(json.loads(raw),ensure_ascii=False,indent=2)
        except Exception: pass
    elif suffix=='.csv':
        try: raw='\n'.join(' | '.join(row) for row in csv.reader(raw.splitlines()))
        except Exception: pass
    raw=_clean(raw)
    if not raw: raise ValueError('Document contains no readable text')
    return [(1,raw,{'ocr':False})]


def read_document(path: str):
    return '\n'.join(text for _,text,*_ in read_document_pages(path))


def ingest_file(path, store, metadata_base=None):
    p=Path(path); metadata_base=dict(metadata_base or {}); results=[]; global_chunk=0
    for entry in read_document_pages(path):
        page,text,extra = entry if len(entry)==3 else (*entry,{})
        chunks=chunk_text(text)
        for chunk_index,chunk in enumerate(chunks):
            meta={**metadata_base,'source':metadata_base.get('source',p.name),'page':page,
                  'chunk':global_chunk,'page_chunk':chunk_index,'ocr':bool(extra.get('ocr',False))}
            results.append(store.add(chunk,meta)); global_chunk+=1
    return results
